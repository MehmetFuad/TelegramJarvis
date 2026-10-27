"""create_document tool — Word (.docx) belgesi olusturma.

Icerigi (siir, metin vb.) Gemini uretir; bu tool sadece docx'e doker.
Kayit: data/documents/<basliktan-guvenli-ad>.docx
Ileride PDF/Excel/PPTX icin ayni klasorde benzer fonksiyonlar acilabilir.
"""
from __future__ import annotations

import re

from config import settings

TOOL_SCHEMA = {
    "name": "create_document",
    "description": (
        "Verilen baslik ve icerikle bir Word (.docx) belgesi olusturur. "
        "Icerigi once kendin uret (ornegin istenen siirleri/metni yaz), sonra "
        "bu tool ile dosyaya dok. Paragraflari cift satir sonu (\\n\\n) ile ayir."
    ),
    "parameters": {
        "type": "object",
        "properties": {
            "title": {
                "type": "string",
                "description": "Belge basligi (Belge icinde Baslik 1 olur, dosya adinda da kullanilir).",
            },
            "content": {
                "type": "string",
                "description": "Belge icerigi; paragraflar \\n\\n ile ayrilir.",
            },
        },
        "required": ["title", "content"],
    },
}

_UNSAFE_CHARS = re.compile(r'[\\/:*?"<>|]')


def safe_filename(title: str) -> str:
    """Basligi dosya adi icin guvenli hale getirir."""
    name = _UNSAFE_CHARS.sub("", (title or "").strip())
    name = re.sub(r"\s+", "_", name).strip("_")
    return name[:60] or "belge"


def create_document(title: str, content: str) -> dict:
    """Docx belgesi olusturur; sonucu standart tool dict'i olarak doner."""
    try:
        from docx import Document
    except ImportError:
        return {"success": False, "error": "python-docx paketi kurulu degil ('pip install python-docx')."}
    title = (title or "").strip()
    content = (content or "").strip()
    if not title or not content:
        return {"success": False, "error": "Baslik ve icerik bos olmamali."}
    try:
        out_dir = settings.data_dir / "documents"
        out_dir.mkdir(parents=True, exist_ok=True)
        path = out_dir / f"{safe_filename(title)}.docx"

        doc = Document()
        doc.add_heading(title, level=1)
        for paragraph in content.split("\n\n"):
            paragraph = paragraph.strip()
            if paragraph:
                doc.add_paragraph(paragraph)
        doc.save(str(path))
        return {"success": True, "file_to_send": str(path)}
    except Exception as exc:  # noqa: BLE001
        return {"success": False, "error": f"Belge olusturulamadi: {exc}"}
